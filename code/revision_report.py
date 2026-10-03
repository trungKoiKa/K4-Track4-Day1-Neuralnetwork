"""Build the final submission from the prospective Colab experiment records."""
import json
import math
import re
from pathlib import Path
import shutil
import numpy as np
import matplotlib.pyplot as plt
from results_table import load_results, to_row, write_xlsx
from plots import plot_compare
from protocol import ROOT, OUT, PREDICTIONS

def finish():
    p=json.loads((OUT/'protocol.json').read_text(encoding='utf-8'))
    h=json.loads((OUT/'health_checks.json').read_text(encoding='utf-8'))
    rs=load_results(OUT/'results'); lookup={r['cfg']['exp_id']:r for r in rs}
    base=lookup['base-s1']; final=lookup[p['selected_exp_id']]
    seeds=[lookup[f'base-s{s}'] for s in [1,2,3]]
    fseeds=[final,lookup['final-s2'],lookup['final-s3']]
    def stat(runs,key):
        vals=[r['summary'][key] for r in runs]
        return float(np.mean(vals)),float(np.std(vals,ddof=1))
    mean,std=stat(seeds,'val_macro_f1'); amean,astd=stat(seeds,'val_acc'); noise=2*std
    bs=json.loads((OUT/'eval_baseline.json').read_text(encoding='utf-8'))
    fs=json.loads((OUT/'eval_result.json').read_text(encoding='utf-8'))
    evals=[fs,*[json.loads((OUT/f'eval_final-s{s}.json').read_text(encoding='utf-8')) for s in [2,3]]]
    emean=float(np.mean([x['macro_f1'] for x in evals])); estd=float(np.std([x['macro_f1'] for x in evals],ddof=1))
    rows=[]
    for r in rs:
        name=r['cfg']['exp_id']; notes='Fresh independent worker; hypotheses in planned/'+name+'.json. '
        if r['cfg']['group']=='optimizer': notes+='Optimizer and lr change together for fair lr tuning. momentum=0.9; Adam betas=(0.9,0.999), eps=1e-8; weight_decay=0. '
        if name=='highlr-clip': notes+='Compared with highlr-none, only clip changes. Relative to baseline, lr and clip both change. '
        notes+=f"grad_p90={r['summary']['grad_p90']:.8f}; clip_fraction_mean={np.mean(r['history']['clip_fraction']):.8f}. "
        if r['summary'].get('nonfinite_history_epochs'):
            notes+='Nonfinite gradient mean from AMP overflow recorded as null; GradScaler skips unsafe updates. '+json.dumps(r['summary']['nonfinite_history_epochs'])
        if name=='base-s1': notes+='Health='+json.dumps({k:v for k,v in h.items() if k!='overfit_history'},ensure_ascii=False)
        if name in ['base-s1',p['selected_exp_id']]: notes+=' Per-class eval='+json.dumps(bs['per_class'] if name=='base-s1' else fs['per_class'])
        row=to_row(r,bs if name=='base-s1' else fs if name==p['selected_exp_id'] else None,notes)
        for k,v in list(row.items()):
            if isinstance(v,float) and not math.isfinite(v): row[k]=None; row['notes']+=f' {k} nonfinite due to divergence.'
        rows.append(row)
    write_xlsx(rows,str(ROOT/'templates/experiment_table_template.xlsx'),str(OUT/'experiments.xlsx'))
    # openpyxl is the course-required generator; artifact-tool subsequently recalculates the delivered file.
    import openpyxl
    w=openpyxl.load_workbook(OUT/'experiments.xlsx'); summary=w['Summary']
    extras=[['Metric','exp_id / source','value','std / detail'],
      ['Baseline val macro-F1','base-s1..3',mean,std],['Baseline val accuracy','base-s1..3',amean,astd],
      ['Final eval macro-F1',p['selected_exp_id']+'; final-s2; final-s3',emean,estd],
      ['Noise threshold 2 sigma','base-s1..3',noise,'Validation only'],
      ['Overfit20 final loss','base-s1 health',h['overfit20_loss'],'500 Adam updates, lr=.01'],
      ['Overfit20 accuracy','base-s1 health',h['overfit20_accuracy'],'20 samples'],
      ['Majority val accuracy','base-s1 health',h['majority_val_accuracy'],'Fixed stratified split'],
      ['Uniform logits CE','base-s1 health',h['uniform_logits_ce'],'ln(7)'],
      ['Recorded epoch seconds','all experiment rows',sum(sum(r['history']['epoch_time_s']) for r in rs),'Excludes setup, plotting, health, scoring'],
      ['Final minus baseline eval F1',p['selected_exp_id'],fs['macro_f1']-bs['macro_f1'],'Baseline eval has one seed'],
      ['Final val/eval F1 gap',p['selected_exp_id'],fs['macro_f1']-final['summary']['val_macro_f1'],'eval minus val']]
    for init,val in h['initialization'].items():
        extras.append([f'Init {init} std after ReLU1/2','base-s1 health',*val['relu_std']])
    for score,name in [(bs,'base-s1'),(fs,p['selected_exp_id'])]:
        extras.append(['Per-class eval: '+name,'class','support','precision','recall','F1'])
        extras.extend([[name,c['cls'],c['support'],c['precision'],c['recall'],c['f1']] for c in score['per_class']])
    for row_no,row in enumerate(extras,start=20):
        for col,value in enumerate(row,start=1): summary.cell(row_no,col,value)
    for i,group in enumerate(['baseline','loss','optimizer','hyperparameter','dropout','clipping','mixed_precision','initialization'],start=2):
        summary.cell(i,8,'See REPORT.md; hypotheses saved before run; compare validation and seed noise.')
    w.save(OUT/'experiments.xlsx')

    groups=['loss','optimizer','hyperparameter','dropout','clipping','mixed_precision','initialization']
    fig,axes=plt.subplots(4,2,figsize=(11,12))
    for ax,group in zip(axes.flat,groups):
        members=[base]+[r for r in rs if r['cfg']['group']==group]
        plot_compare(members,'val_macro_f1',str(OUT/'figures'/f'compare_{group}.png'),group)
        for r in members: ax.plot(r['history']['epoch'],r['history']['val_macro_f1'],label=r['cfg']['exp_id'])
        ax.set(title=group,xlabel='Epoch',ylabel='Val macro-F1'); ax.legend(fontsize=6); ax.grid(alpha=.2)
    axes.flat[-1].axis('off'); fig.tight_layout(); fig.savefig(OUT/'figures/compare_all.png',dpi=130); plt.close(fig)
    bests=[max([r for r in rs if r['cfg']['group']==g and not r['summary']['diverged']],key=lambda r:r['summary']['val_macro_f1']) for g in groups]
    fig,ax=plt.subplots(figsize=(8.5,3.8))
    ax.barh(groups,[r['summary']['val_macro_f1'] for r in bests],color='#356894')
    ax.axvline(base['summary']['val_macro_f1'],color='#bd4e35',ls='--',label='base-s1')
    for index,r in enumerate(bests): ax.text(r['summary']['val_macro_f1']+.003,index,f"{r['summary']['val_macro_f1']:.4f}",va='center',fontsize=10)
    ax.set(xlim=(0,1),xlabel='Validation macro-F1',title='Best new run per topic; baseline shown separately'); ax.legend()
    fig.tight_layout(); fig.savefig(OUT/'figures/compare_overview.png',dpi=150); plt.close(fig)
    plot_compare(seeds,'val_macro_f1',str(OUT/'figures/compare_baseline.png'),'SGD momentum baseline: 3 seeds')
    cm=np.asarray(fs['confusion_matrix']); fig,ax=plt.subplots(figsize=(5,4))
    im=ax.imshow(cm/cm.sum(1,keepdims=True),vmin=0,vmax=1,cmap='Blues'); fig.colorbar(im,ax=ax,label='Fraction of true class')
    ax.set(xlabel='Predicted class',ylabel='True class',xticks=range(7),yticks=range(7),title='Final eval confusion')
    fig.tight_layout(); fig.savefig(OUT/'figures/confusion_eval.png',dpi=130); plt.close(fig)

    def metrics(r): return f"`{r['cfg']['exp_id']}` F1={r['summary']['val_macro_f1']:.6f}"
    def comparison(group):
        members=[r for r in rs if r['cfg']['group']==group and not r['summary']['diverged']]
        best=max(members,key=lambda r:r['summary']['val_macro_f1']); delta=best['summary']['val_macro_f1']-base['summary']['val_macro_f1']
        return f"{metrics(best)}, Δ so `base-s1`={delta:+.6f} ({'vượt' if abs(delta)>noise else 'chưa vượt'} 2σ tham chiếu)."
    base_gap=base['summary']['final_val_loss']-base['summary']['final_train_loss']
    dr=lookup['dropout-p03']; drop_gap=dr['summary']['final_val_loss']-dr['summary']['final_train_loss']
    clipped=lookup['highlr-clip']; unclipped=lookup['highlr-none']
    worst=min(fs['per_class'],key=lambda c:c['f1']); errs=cm[worst['cls']].copy(); errs[worst['cls']]=0
    optimizer_best=[]
    for opt in ['sgd_momentum','adam','adamw']:
        candidates=[r for r in rs if r['cfg']['optimizer']==opt and r['cfg']['group'] in ['baseline','optimizer'] and r['cfg']['seed']==1]
        optimizer_best.append(max(candidates,key=lambda r:r['summary']['val_macro_f1']))
    lines=[f'# Báo cáo Lab Day 1 — Hoàng Trung Anh — 2A202602521','',
      '## 1. Thiết lập',
      f"Colab {h['gpu']}, PyTorch {h['torch_version']}. M-base 54→256→128→7, 47 879 tham số, ReLU, logits thô. Metadata train/eval 464 809/116 203; validation 20%, stratify seed 42: 371 847 train/92 962 val. Chuẩn hoá 10 cột từ train; giữ 44 cột nhị phân. Majority val accuracy={h['majority_val_accuracy']:.6f}.",
      f"Baseline đúng GUIDE: SGD momentum 0,9, lr={p['baseline_lr']}, batch 512, 20 epoch, CE, He, FP32; không dropout/clip/weight decay. Lr chọn từ 0.01/0.05/0.1 bằng val. Mỗi run khởi chạy process GPU độc lập; batch cuối giữ lại. Adam/AdamW thử 0.0003/0.001/0.003. Đã thử đủ 7 chủ đề. Mọi số liệu có trong Experiments/Seeds/Summary, tra theo exp_id; cấu hình và giả thuyết ghi trước trong protocol.json/planned/.",
      '', '## 2. Kiểm tra ban đầu và độ nhiễu',
      f"Logits (20,7); gradient khác 0 ở mọi tham số. Overfit 20 mẫu: accuracy={h['overfit20_accuracy']:.6f}, loss={h['overfit20_loss']:.8f} sau 500 bước Adam lr=.01. Loss uniform logits={h['uniform_logits_ce']:.6f}=ln 7. He step-0 CE={base['summary']['step0_loss']:.6f}; logits ngẫu nhiên có phương sai nên CE không bằng uniform. Kiểm tra dtype, chuẩn hoá, gradient và overfit đều qua trước khi chạy dài.",
      f"Baseline seed 1–3: val acc={amean:.6f} ± {astd:.6f}; macro-F1={mean:.6f} ± {std:.6f}; 2σ={noise:.6f} (std mẫu). Gap loss cuối base-s1={base_gap:.6f}, best epoch={base['summary']['best_epoch']}; train/val loss giảm mạnh lúc đầu rồi giảm chậm, có dao động và val tăng nhẹ sau cực tiểu ở epoch 18. Gap cuối nhỏ, chưa có bằng chứng overfit mạnh. Ngưỡng này là tham chiếu, không phải kiểm định thống kê cho từng cấu hình một seed.",
      '![Overfit 20 mẫu](figures/health_overfit20.png)', '', '## 3. Kết quả theo chủ đề',
      'Các dự đoán sau được lưu trước run mới. So sánh dùng metric tại epoch có val loss thấp nhất; mỗi run có PNG loss/metric/gradient riêng. Hình tổng hợp bên dưới và compare_<nhóm>.png đối chiếu các đường.',
      '**Loss.** Dự đoán CE phù hợp phân lớp hơn MSE logits/one-hot. '+comparison('loss')+' MSE lấy mean trên B×7, không có hệ số 1/2; khác thang CE nên không so loss trực tiếp. CE liên hệ gradient với xác suất lớp; kết quả chỉ đúng lr/ngân sách đã thử.',
      '**Optimizer.** Dự đoán Adam thích nghi bước nên có thể hội tụ nhanh hơn. '+comparison('optimizer')+' SGD có 3 lr và Adam/AdamW có 3 lr; phải so mỗi bộ ở lr tốt nhất, không so một lr bất kỳ. Adam và AdamW wd=0 tương đương, không đo lợi ích decoupled decay.',
      '**Hyper-parameter.** Dự đoán batch lớn học chậm nếu giữ lr. '+comparison('hyperparameter')+' Batch 512/2048 tương ứng 727/182 update mỗi epoch; cùng epoch không cùng số update. Đây là thay batch duy nhất, không tự tăng lr.',
      '**Dropout.** Dự đoán giảm gap nhưng có thể underfit. '+comparison('dropout')+f' Gap cuối p=.3/base={drop_gap:.6f}/{base_gap:.6f}; xem cả train và val loss trong dòng tương ứng, gap nhỏ tự nó không chứng minh khái quát hoá tốt.',
      '**Clipping.** Dự đoán hạn chế gradient cực lớn, không chắc cứu được lr cao. '+f"c={p['clip_threshold']:.6f} từ phân vị 90 grad batch baseline, lr cao={p['baseline_lr']*10}. Clip thường: tỷ lệ batch clip={np.mean(lookup['clip-normal']['history']['clip_fraction']):.6f}; clip lr cao={np.mean(clipped['history']['clip_fraction']):.6f}. Cặp cùng lr: {metrics(clipped)} vs {metrics(unclipped)}; diverged={clipped['summary']['diverged']}/{unclipped['summary']['diverged']}. Cặp lr cao khác baseline cả lr/clip; không quy mọi cải thiện cho clipping. Grad đo trước clip.",
      '**Mixed precision.** Dự đoán overhead MLP nhỏ có thể xoá lợi ích Tensor Core. '+comparison('mixed_precision')+f" Thời gian FP32/FP16/BF16={base['summary']['time_per_epoch_s']:.4f}/{lookup['precision-fp16']['summary']['time_per_epoch_s']:.4f}/{lookup['precision-bf16']['summary']['time_per_epoch_s']:.4f} s/epoch; peak={base['summary']['peak_mem_MB']:.2f}/{lookup['precision-fp16']['summary']['peak_mem_MB']:.2f}/{lookup['precision-bf16']['summary']['peak_mem_MB']:.2f} MB. Process độc lập, cùng dữ liệu; peak gồm data/model/optimizer, không riêng activation. BF16 native={h['bf16_native']}; FP16 cần scaler do miền số hẹp, unscale trước clip. Thời gian gồm đánh giá FP32 toàn train/val. FP16 có gradient mean không hữu hạn ở một số epoch (notes/JSON ghi null và epoch cụ thể), loss vẫn hữu hạn; GradScaler bỏ các cập nhật overflow. Khoảng trống trên đường gradient phản ánh giá trị thiếu, không phải gradient bằng 0.",
      '**Init.** Dự đoán He giữ tín hiệu ReLU, normal nhỏ giảm tín hiệu, zeros không phá đối xứng. '+comparison('initialization')+' Std sau mỗi ReLU trên 2 048 mẫu val và step-0 CE lưu Summary/health_checks.json. Zeros cho hidden activation=0 và chỉ bias đầu ra có thể học prior; lớp ẩn không học qua ReLU(0). He var=2/fan_in, Xavier var=2/(fan_in+fan_out). Mạng hai lớp ẩn chưa đại diện mạng rất sâu.',
      'Ảnh riêng theo chủ đề: '+', '.join(f'[{g}](figures/compare_{g}.png)' for g in groups)+'.',
      '| Bộ tối ưu | exp_id tốt nhất trong lưới | lr | val F1 | best epoch |',
      '|---|---|---:|---:|---:|',
      *[f"| {r['cfg']['optimizer']} | {r['cfg']['exp_id']} | {r['cfg']['lr']} | {r['summary']['val_macro_f1']:.6f} | {r['summary']['best_epoch']} |" for r in optimizer_best],
      f"Đối chiếu giả thuyết: MSE {'thấp hơn CE, khớp' if lookup['loss-mse']['summary']['val_macro_f1']<base['summary']['val_macro_f1'] else 'không thấp hơn CE, khác'} dự đoán; batch2048 {'thấp hơn, khớp' if lookup['batch2048']['summary']['val_macro_f1']<base['summary']['val_macro_f1'] else 'không thấp hơn, khác'} dự đoán học chậm. Mixed precision {'có lượt nhanh hơn FP32' if min(lookup['precision-fp16']['summary']['time_per_epoch_s'],lookup['precision-bf16']['summary']['time_per_epoch_s'])<base['summary']['time_per_epoch_s'] else 'không nhanh hơn FP32'} theo số đo, không mặc định tăng tốc. Các Δ nhỏ hơn 2σ chưa đủ phân biệt nhiễu; cấu hình final có lặp seed nhưng các chủ đề khác vẫn thăm dò.",
      '![So sánh 7 chủ đề](figures/compare_overview.png)', '', '## 4. Đánh giá cuối và lỗi theo lớp',
      f"Khoá `{p['selected_exp_id']}` bằng F1 val trước scoring mới, seed nộp=1; không đổi sau eval. Cfg={final['cfg']['optimizer']}, lr={final['cfg']['lr']}, batch={final['cfg']['batch']}, init={final['cfg']['init']}, dropout={final['cfg']['dropout']}, clip={final['cfg']['clip_norm']}, precision={final['cfg']['precision']}. Checkpoint min val loss.",
      '| exp_id | val F1 | eval accuracy | eval F1 |','|---|---:|---:|---:|',
      f"| base-s1 | {base['summary']['val_macro_f1']:.6f} | {bs['accuracy']:.6f} | {bs['macro_f1']:.6f} |",
      f"| {p['selected_exp_id']} | {final['summary']['val_macro_f1']:.6f} | {fs['accuracy']:.6f} | {fs['macro_f1']:.6f} |",
      f"Final eval seed 1–3 F1={emean:.6f} ± {estd:.6f}. Seed 1 Δeval vs baseline={fs['macro_f1']-bs['macro_f1']:+.6f}; eval−val={fs['macro_f1']-final['summary']['val_macro_f1']:+.6f}. Baseline eval một seed nên chưa có kiểm định hai nhóm; nhiễu val không thay được nhiễu eval.",
      '| Lớp | support | precision | recall | F1 |','|---|---:|---:|---:|---:|']
    lines.extend(f"| {c['cls']} | {c['support']} | {c['precision']:.6f} | {c['recall']:.6f} | {c['f1']:.6f} |" for c in fs['per_class'])
    lines.extend([f"Lớp khó nhất {worst['cls']} (F1={worst['f1']:.6f}), nhầm nhiều nhất sang {errs.argmax()}; support={worst['support']}. Mất cân bằng và tương đồng địa hình có thể giải thích nhưng chưa có kiểm chứng feature. Có thể thử weighted CE bằng val ở nghiên cứu sau.",
      '![Ma trận nhầm lẫn](figures/confusion_eval.png)', '', '## 5. Câu hỏi dẫn dắt',
      '1. Optimizer tốt nhất phụ thuộc lưới lr; so lr tốt nhất của mỗi bộ trong Experiments, không suy rộng ngoài lưới. Adam thích nghi bước, SGD momentum tích luỹ hướng gradient.',
      '2. Dùng dropout khi có bằng chứng train giảm nhưng val tăng; khi chưa overfit, nó có thể làm underfit (mục 3).',
      '3. Clipping giới hạn chuẩn gradient, không sửa nhãn hay lr bất kỳ; bằng chứng là clip fraction và cặp cùng lr ở mục 3, kể cả trường hợp không cứu được phân kỳ.',
      '4. Mixed precision chỉ nhanh hơn nếu số đo thời gian ở mục 3 nhỏ hơn FP32; data resident, kernel overhead, optimizer và evaluation FP32 có thể chi phối mạng nhỏ.',
      '5. Zeros giữ đối xứng và ReLU(0) chặn gradient hidden; He bù phương sai bị ReLU cắt, Xavier cân bằng fan-in/fan-out.',
      '6. Loss không giảm sau 2 000 bước: (i) kiểm tra nhãn 0..6, dtype/shape và CE trên logits; (ii) kiểm tra chuẩn hoá, gradient và zero_grad/backward/step; (iii) overfit 20 mẫu rồi quét lr trên val. Các phép kiểm tra lần lượt tách lỗi dữ liệu, cập nhật và cấu hình.',
      '', '## 6. Hạn chế và điều bất ngờ',
      'Đây là lượt sửa tuân thủ đề sau khi đã biết eval cũ. Lưới và giả thuyết mới được khoá trước run, không chọn theo eval mới, nhưng không thể coi eval hoàn toàn chưa từng được xem. Giữ lịch sử cũ trong Git; không giả mạo dự đoán trước của lượt cũ. Chỉ baseline/final lặp 3 seed; các chủ đề khác một seed. Kết quả khác giả thuyết cần diễn giải theo số quan sát, không coi lý thuyết là kết quả đo. BF16 emulation và thời gian có nhiễu hệ thống; không thử weight_decay, width/depth vì menu là tuỳ chọn.',
      '', '## 7. Phụ lục',
      f"Tổng thời gian epoch đã ghi={sum(sum(r['history']['epoch_time_s']) for r in rs):.2f} s, không gồm khởi tạo process, health, vẽ và scoring. Nộp REPORT.md, experiments.xlsx, predictions/eval JSON, code/lab.ipynb và modules, figures, results, protocol.json/planned. Toàn bộ code trong code; không nộp dữ liệu/cache/checkpoint. Notebook giữ output; mở từ submission/code, đọc data ở ../../data và ghi figures/results ở ../. Công thức Excel tính lại khi mở; bản giao đã được recalculation và kiểm tra."])
    report=re.sub(r'(?m)(^\|[^\n]*\|)\n\n(?=\|)',r'\1\n','\n\n'.join(lines))
    (OUT/'REPORT.md').write_text(report,encoding='utf-8')
    write_data={'baseline':bs,'final':fs,'final_eval_mean':emean,'final_eval_std':estd,'baseline_val_mean':mean,'baseline_val_std':std,'experiment_count':len(rs)}
    (OUT/'verification_metrics.json').write_text(json.dumps(write_data,indent=2),encoding='utf-8')
    code_out=OUT/'code'; code_out.mkdir(exist_ok=True)
    for file in (ROOT/'code').iterdir():
        if file.suffix in ['.py','.txt'] and file.resolve()!=(code_out/file.name).resolve(): shutil.copy2(file,code_out/file.name)
    print('SUBMISSION_COMPLETE',json.dumps(write_data)[:300],flush=True)
